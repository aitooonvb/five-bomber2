#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
ファイブボンバー風 ネットワークゲーム【インターネット版サーバ】

  /       … 子機（プレイヤー）用ページ
  /host   … 親機（先生・司会）用ページ
  /ws     … WebSocket（ゲームの通信）

問題は同じフォルダの questions.xlsx から読み込みます（部屋を作るとき／「問題リセット」のとき）。
起動:  python server.py      （ポートは環境変数 PORT、なければ 8080）
"""
import asyncio
import json
import os
import random
import time
import unicodedata
from pathlib import Path

from aiohttp import WSMsgType, web
from openpyxl import load_workbook

BASE = Path(__file__).parent
STATIC = BASE / "static"
QUESTIONS_FILE = Path(os.environ.get("QUESTIONS_FILE", BASE / "questions.xlsx"))
N = 5  # 解答の文字数 = 子機の台数


def load_questions():
    """questions.xlsx から (問題, 答え) のリストを読み込む。答えが5文字でない行は除外。"""
    if not QUESTIONS_FILE.exists():
        raise ValueError(f"問題ファイル {QUESTIONS_FILE.name} が見つかりません（server.py と同じフォルダに置いてください）")
    wb = load_workbook(QUESTIONS_FILE, read_only=True, data_only=True)
    ws = wb["問題"] if "問題" in wb.sheetnames else wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    qcol, acol, start = 0, 1, 0          # 見出し行に「問題」「答え」があればその列を使う
    if rows:
        head = [str(c).strip() if c is not None else "" for c in rows[0]]
        if "問題" in head and "答え" in head:
            qcol, acol, start = head.index("問題"), head.index("答え"), 1

    out, skipped = [], []
    for i, row in enumerate(rows[start:], start=start + 1):
        if not row or len(row) <= max(qcol, acol):
            continue
        q, a = row[qcol], row[acol]
        if q is None and a is None:
            continue
        q = "" if q is None else str(q).strip()
        a = "" if a is None else str(a).strip().replace(" ", "").replace("\u3000", "")
        if q and len(a) == N:
            out.append((q, a))
        else:
            skipped.append(i)
    if skipped:
        print(f"[注意] {N}文字の答えになっていない行を除外しました（Excelの行番号）: {skipped}")
    if not out:
        raise ValueError(f"{QUESTIONS_FILE.name} に使える問題がありません（答えは{N}文字にしてください）")
    return out


def normalize(s):
    """全角/半角・大文字小文字・カタカナ/ひらがなの違いを吸収して比較"""
    s = unicodedata.normalize("NFKC", s).strip().lower()
    return "".join(chr(ord(c) - 0x60) if "\u30a1" <= c <= "\u30f6" else c for c in s)


class Room:
    def __init__(self, code, host_key, questions):
        self.code = code
        self.host_key = host_key      # 親機の乗っ取り防止用の秘密キー
        self.host = None
        self.players = {}             # slot(1-5) -> WebSocket
        self.owner = {}               # slot -> cid（再接続の本人確認用）
        self.cids = {}                # cid -> 最後に使った slot
        self.questions = questions
        self.deck = []                # まだ出していない問題の番号（シャッフル済み）
        self.qi = -1                  # 現在の問題番号
        self.round = 0                # 出題のたびに増える（子機の入力欄リセット用）
        self.phase = "lobby"          # lobby / answering / ready / revealed
        self.inputs = [None] * N
        self.ok = None
        self.wrong = [False] * N      # 公開時：間違っている文字の位置
        self.touched = time.time()
        self.reset_deck()

    def reset_deck(self):
        self.deck = list(range(len(self.questions)))
        random.shuffle(self.deck)     # 1周するまで同じ問題は出ない

    def snapshot(self):
        revealed = self.phase == "revealed"
        return {
            "type": "state", "room": self.code, "phase": self.phase, "round": self.round,
            "qno": len(self.questions) - len(self.deck) if self.qi >= 0 else 0,
            "text": self.questions[self.qi][0] if self.qi >= 0 else "",
            "asked": len(self.questions) - len(self.deck),
            "total": len(self.questions),
            "remaining": len(self.deck),
            "connected": [s in self.players for s in range(1, N + 1)],
            "submitted": [c is not None for c in self.inputs],
            # 正解・各自の文字は公開するまで送らない
            "chars": list(self.inputs) if revealed else None,
            "ok": self.ok if revealed else None,
            "wrong": list(self.wrong) if revealed else None,
            "answer": self.questions[self.qi][1] if revealed else None,
        }


rooms = {}


async def safe_send(ws, obj):
    try:
        if ws is not None and not ws.closed:
            await ws.send_str(json.dumps(obj, ensure_ascii=False))
    except Exception:
        pass


async def broadcast(room):
    room.touched = time.time()
    snap = room.snapshot()
    await safe_send(room.host, snap)
    for slot, ws in list(room.players.items()):
        await safe_send(ws, {**snap, "you": slot, "mine": room.inputs[slot - 1]})


async def fail(ws, message):
    await safe_send(ws, {"type": "error", "message": message, "fatal": True})
    await ws.close()


async def ws_handler(request):
    ws = web.WebSocketResponse(heartbeat=20)
    await ws.prepare(request)
    room = role = slot = None

    try:
        async for msg in ws:
            if msg.type != WSMsgType.TEXT:
                continue
            try:
                d = json.loads(msg.data)
            except ValueError:
                continue
            t = d.get("type")

            # ---------- 最初のあいさつ ----------
            if t == "hello" and room is None:
                code = "".join(ch for ch in str(d.get("room", "")).upper() if ch.isalnum())[:8]
                if len(code) < 3:
                    await fail(ws, "部屋コードを入力してください")
                    break

                if d.get("role") == "host":
                    key = str(d.get("key", ""))[:64]
                    r = rooms.get(code)
                    if r is None:
                        try:
                            r = rooms[code] = Room(code, key, load_questions())
                        except ValueError as e:
                            await fail(ws, str(e))
                            break
                    elif r.host_key != key:
                        await fail(ws, "この部屋コードはすでに使われています。別のコードにしてください")
                        break
                    old = r.host
                    r.host = ws
                    room, role = r, "host"
                    if old is not None and old is not ws:
                        try:
                            await old.close()
                        except Exception:
                            pass

                elif d.get("role") == "player":
                    r = rooms.get(code)
                    if r is None:
                        await fail(ws, "その部屋コードの部屋が見つかりません（親機を先に開いてください）")
                        break
                    cid = str(d.get("cid", ""))[:64]
                    prev = r.cids.get(cid)
                    if prev and r.owner.get(prev) == cid:      # 同じ端末の再接続
                        old = r.players.get(prev)
                        r.players[prev] = ws
                        slot = prev
                        if old is not None and old is not ws:
                            try:
                                await old.close()
                            except Exception:
                                pass
                    else:
                        if prev and prev not in r.players:
                            slot = prev
                        else:
                            free = [s for s in range(1, N + 1) if s not in r.players]
                            slot = free[0] if free else None
                        if slot is None:
                            await fail(ws, "この部屋はすでに5人そろっています")
                            break
                        r.players[slot] = ws
                    r.owner[slot] = cid
                    r.cids[cid] = slot
                    room, role = r, "player"
                else:
                    await fail(ws, "不正な接続です")
                    break
                await broadcast(room)

            # ---------- 親機の操作 ----------
            elif room and role == "host":
                if t in ("next", "retry"):
                    if len(room.players) < N:
                        await safe_send(ws, {"type": "error", "message": f"{N}人そろってから出題してください"})
                        continue
                    if room.phase in ("answering", "ready"):
                        continue
                    if t == "retry" and room.qi >= 0:
                        pass                                  # 同じ問題をもう一度
                    elif room.deck:
                        room.qi = room.deck.pop()             # 未出題の問題から1問（重複なし）
                    else:
                        await safe_send(ws, {"type": "error", "message": "すべての問題を出題しました。「問題リセット」で最初からできます"})
                        continue
                    room.round += 1
                    room.phase = "answering"
                    room.inputs = [None] * N
                    room.ok = None
                    await broadcast(room)

                elif t == "judge" and room.phase == "ready":   # 「答えあわせをしますか？」→ はい
                    answer = room.questions[room.qi][1]
                    room.wrong = [normalize(c) != normalize(a) for c, a in zip(room.inputs, answer)]
                    room.ok = not any(room.wrong)
                    room.phase = "revealed"
                    await broadcast(room)

                elif t == "reset_deck" and room.phase in ("lobby", "revealed"):
                    try:
                        room.questions = load_questions()      # Excelを編集していれば最新を読み直す
                    except ValueError as e:
                        await safe_send(ws, {"type": "error", "message": str(e)})
                    room.reset_deck()
                    room.qi, room.phase, room.ok = -1, "lobby", None
                    room.inputs = [None] * N
                    room.round += 1
                    await broadcast(room)

            # ---------- 子機の入力（1回送信したら変更不可） ----------
            elif room and role == "player" and t == "input":
                if room.phase != "answering" or room.inputs[slot - 1] is not None:
                    continue
                ch = str(d.get("char", "")).strip()
                if len(ch) != 1:
                    continue
                room.inputs[slot - 1] = ch
                if all(c is not None for c in room.inputs):
                    room.phase = "ready"                       # 全員そろった → 親機の「はい」待ち
                await broadcast(room)
    finally:
        if room is not None:
            if role == "host" and room.host is ws:
                room.host = None
            elif role == "player" and room.players.get(slot) is ws:
                del room.players[slot]
                room.owner.pop(slot, None)
            await broadcast(room)

    return ws


def find(name):
    """static/ フォルダを探し、無ければ server.py と同じ階層も探す（アップロード時の置き間違い対策）"""
    for d in (STATIC, BASE):
        if (d / name).is_file():
            return d / name
    raise web.HTTPNotFound(text=f"{name} が見つかりません。static フォルダに置いてください")


def page(name):
    async def handler(request):
        return web.FileResponse(find(name))
    return handler


async def health(request):
    return web.Response(text="ok")


async def cleanup_ctx(app):
    """使われなくなった部屋を定期的に削除"""
    async def loop():
        while True:
            await asyncio.sleep(600)
            now = time.time()
            for code, r in list(rooms.items()):
                if r.host is None and not r.players and now - r.touched > 1800:
                    rooms.pop(code, None)
    task = asyncio.create_task(loop())
    yield
    task.cancel()


def make_app():
    app = web.Application()
    app.router.add_get("/", page("client.html"))
    app.router.add_get("/host", page("host.html"))
    app.router.add_get("/ws", ws_handler)
    app.router.add_get("/health", health)
    app.router.add_get("/static/style.css", page("style.css"))
    app.cleanup_ctx.append(cleanup_ctx)
    return app


if __name__ == "__main__":
    port = int(os.environ.get("PORT", 8080))
    try:
        print(f"問題を {len(load_questions())} 問読み込みました（{QUESTIONS_FILE.name}）")
    except ValueError as e:
        print(f"[エラー] {e}")
    missing = [n for n in ("host.html", "client.html", "style.css") if not any((d / n).is_file() for d in (STATIC, BASE))]
    if missing:
        print(f"[警告] 次のファイルが見つかりません（static フォルダに置いてください）: {missing}")
    print(f"起動しました:  親機 http://localhost:{port}/host   子機 http://localhost:{port}/")
    web.run_app(make_app(), host="0.0.0.0", port=port)
